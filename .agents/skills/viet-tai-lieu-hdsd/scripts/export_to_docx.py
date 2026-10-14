"""
Tiện ích xuất bản tài liệu Hướng dẫn sử dụng sang định dạng Microsoft Word (.docx).
Yêu cầu: pip install python-docx
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_hdsd_docx(output_path, title, subtitle, metadata_list, sections_data, trouble_data):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1.0)

    # Base styling
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10.5)

    COLOR_PRIMARY = RGBColor(0x00, 0x3e, 0xb3)
    COLOR_DARK = RGBColor(0x14, 0x14, 0x14)
    COLOR_MUTED = RGBColor(0x59, 0x59, 0x59)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title)
    r_title.font.size = Pt(19)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.size = Pt(12)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_DARK

    doc.save(output_path)
    print(f"Đã tạo file Word: {output_path}")
