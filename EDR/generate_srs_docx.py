import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def add_formatted_text(paragraph, text, default_color=None, default_size=9.5, is_bold=False, is_italic=False, font_name="Segoe UI"):
    # Split by bold **...**, inline code `...`, italic *...*
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = font_name
        run.font.size = Pt(default_size)
        if default_color:
            run.font.color.rgb = default_color

        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
            run.italic = is_italic
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(default_size - 0.5)
            run.font.color.rgb = RGBColor(194, 24, 91)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.text = token[1:-1]
            run.italic = True
            run.bold = is_bold
        else:
            run.text = token
            run.bold = is_bold
            run.italic = is_italic

def create_docx(md_path, docx_path):
    doc = docx.Document()

    # Page setup - A4, standard professional margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.8)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("HỆ THỐNG EDR | ĐẶC TẢ TÍNH NĂNG QUẢN LÝ IOC & CHIA SẺ DỮ LIỆU")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = RGBColor(148, 163, 184)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun1 = fp.add_run("BẢO MẬT NỘI BỘ - PHÂN HỆ THREAT INTELLIGENCE")
        frun1.font.name = "Segoe UI"
        frun1.font.size = Pt(8)
        frun1.font.color.rgb = RGBColor(148, 163, 184)

    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    lines = md_text.splitlines()

    mermaid_idx = 0
    in_code_block = False
    code_block_lines = []
    code_block_lang = ""

    in_table = False
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        rows_data = []
        for tl in tbl_lines:
            cells = [c.strip() for c in tl.strip().strip('|').split('|')]
            if all(set(c).issubset(set('-: ')) for c in cells):
                continue
            rows_data.append(cells)
        
        if not rows_data:
            return

        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        for r_idx, row in enumerate(rows_data):
            for c_idx in range(num_cols):
                cell_text = row[c_idx] if c_idx < len(row) else ""
                cell = table.cell(r_idx, c_idx)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=70, bottom=70, left=110, right=110)

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.1

                if r_idx == 0:
                    set_cell_background(cell, "1E3A8A")
                    add_formatted_text(p, cell_text, default_color=RGBColor(255, 255, 255), default_size=9, is_bold=True)
                else:
                    bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg_color)
                    add_formatted_text(p, cell_text, default_color=RGBColor(30, 41, 59), default_size=9)
        
        # Add spacing after table
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(6)

    i = 0
    diagram_files = [
        ("diagrams/diagram_1.png", "Hình 1: Sơ đồ Quan hệ Thực thể (ERD) Quản lý IOC & Cảnh báo liên quan"),
        ("diagrams/diagram_2.png", "Hình 2: Sơ đồ Luồng Nghiệp vụ Thu thập, Khử trùng và Chia sẻ Dữ liệu")
    ]

    while i < len(lines):
        line = lines[i]

        # Check code block start / end
        if line.startswith('```'):
            if in_code_block:
                in_code_block = False
                # If was mermaid
                if code_block_lang == 'mermaid':
                    diag_path = None
                    caption_text = ""
                    if mermaid_idx < len(diagram_files):
                        diag_rel, caption_text = diagram_files[mermaid_idx]
                        diag_path = os.path.normpath(os.path.join(os.path.dirname(md_path), diag_rel))
                        mermaid_idx += 1

                    if diag_path and os.path.exists(diag_path):
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_before = Pt(10)
                        p_img.paragraph_format.space_after = Pt(3)
                        p_img.add_run().add_picture(diag_path, width=Inches(6.2))

                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_before = Pt(1)
                        p_cap.paragraph_format.space_after = Pt(10)
                        c_run = p_cap.add_run(caption_text)
                        c_run.font.name = "Segoe UI"
                        c_run.font.size = Pt(8.5)
                        c_run.font.italic = True
                        c_run.font.color.rgb = RGBColor(100, 116, 139)
                else:
                    # Render code snippet box
                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    cell = tbl.cell(0, 0)
                    set_cell_background(cell, "F1F5F9")
                    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                    set_table_borders(tbl, color="CBD5E1", sz="4")

                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(1)
                    cp.paragraph_format.space_after = Pt(1)
                    cp.paragraph_format.line_spacing = 1.05

                    code_content = "\n".join(code_block_lines)
                    c_run = cp.add_run(code_content)
                    c_run.font.name = "Consolas"
                    c_run.font.size = Pt(8.5)
                    c_run.font.color.rgb = RGBColor(30, 41, 59)

                    p_after = doc.add_paragraph()
                    p_after.paragraph_format.space_before = Pt(0)
                    p_after.paragraph_format.space_after = Pt(6)

                code_block_lines = []
                code_block_lang = ""
            else:
                in_code_block = True
                code_block_lang = line[3:].strip().lower()
                code_block_lines = []
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Check Markdown Image: ![Caption](path)
        img_match = re.match(r'^\s*!\[(.*?)\]\((.*?)\)\s*$', line)
        if img_match:
            caption = img_match.group(1).strip()
            rel_path = img_match.group(2).strip()
            base_dir = os.path.dirname(md_path)
            full_img_path = os.path.normpath(os.path.join(base_dir, rel_path))
            
            if os.path.exists(full_img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(3)
                
                p_img.add_run().add_picture(full_img_path, width=Inches(6.2))

                if caption:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(1)
                    p_cap.paragraph_format.space_after = Pt(10)
                    c_run = p_cap.add_run(caption)
                    c_run.font.name = "Segoe UI"
                    c_run.font.size = Pt(8.5)
                    c_run.font.italic = True
                    c_run.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        # Check table
        if line.strip().startswith('|') and line.strip().endswith('|'):
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_lines)
                table_lines = []

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if line.strip() == '---':
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, line[2:].strip(), default_color=RGBColor(30, 58, 138), default_size=15, is_bold=True)
            i += 1
            continue

        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, line[3:].strip(), default_color=RGBColor(30, 58, 138), default_size=13, is_bold=True)
            i += 1
            continue

        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, line[4:].strip(), default_color=RGBColor(37, 99, 235), default_size=11.5, is_bold=True)
            i += 1
            continue

        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, line[5:].strip(), default_color=RGBColor(30, 41, 59), default_size=10.5, is_bold=True)
            i += 1
            continue

        # Lists & Bullets
        list_match = re.match(r'^\s*([0-9]+\.|\-|\*)\s+(.*)', line)
        if list_match:
            bullet_char = list_match.group(1)
            content = list_match.group(2)
            indent_level = len(line) - len(line.lstrip())
            
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.2 + (indent_level // 2) * 0.15)
            
            if bullet_char in ['-', '*']:
                if content.startswith('[x]'):
                    bullet_run = p.add_run("✓ ")
                    bullet_run.bold = True
                    bullet_run.font.color.rgb = RGBColor(16, 185, 129)
                    content = content[3:].strip()
                elif content.startswith('[ ]'):
                    bullet_run = p.add_run("☐ ")
                    bullet_run.font.color.rgb = RGBColor(100, 116, 139)
                    content = content[3:].strip()
                else:
                    bullet_run = p.add_run("• ")
                    bullet_run.bold = True
                    bullet_run.font.color.rgb = RGBColor(37, 99, 235)
            else:
                bullet_run = p.add_run(f"{bullet_char} ")
                bullet_run.bold = True
                bullet_run.font.color.rgb = RGBColor(30, 58, 138)
            
            bullet_run.font.name = "Segoe UI"
            bullet_run.font.size = Pt(9.5)
            add_formatted_text(p, content, default_color=RGBColor(51, 65, 85), default_size=9.5)
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_text(p, line, default_color=RGBColor(51, 65, 85), default_size=9.5)
        i += 1

    if in_table:
        flush_table(table_lines)

    doc.save(docx_path)
    print(f"DOCX created successfully at: {docx_path}")

if __name__ == "__main__":
    md_file = r"E:\work1\EDR\SRS_Quan_Ly_IOC_EDR.md"
    docx_file = r"E:\work1\EDR\SRS_Quan_Ly_IOC_EDR.docx"
    create_docx(md_file, docx_file)
